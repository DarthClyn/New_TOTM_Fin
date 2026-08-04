import io
import base64
import openpyxl

class ConsolidationEngine:
    def __init__(self):
        # Master COA dictionary for matching known accounts
        self.coa_mapping = {
            '1-1300': ('12150302', 'Bank Mandiri (IDR) 662-7'),
            'Bank Account': ('12150202', 'BANK ACCOUNT'),
            'Kas Kecil': ('12150201', 'PETTY CASH'),
            'Petty Cash': ('12150201', 'PETTY CASH'),
            'Amount Due from Subsidiary': ('13150000', 'AMOUNT DUE FROM SUBSIDIARY'),
            'Amount Due to Related Party': ('23150000', 'AMOUNT DUE TO RELATED PARTY'),
            'Sales Revenue': ('41000000', 'SALES REVENUE'),
            'Pendapatan': ('41000000', 'SALES REVENUE'),
            'Staff Salaries': ('71010000', 'STAFF SALARIES'),
            'Gaji': ('71010000', 'STAFF SALARIES'),
            'Salary/Bonus': ('71010000', 'STAFF SALARIES'),
            '6-1200': ('71010000', 'STAFF SALARIES'),
            'ICICI Bank A/C -059805005246': ('12150301', 'ICICI Bank A/C -059805005246'),
            '1-6100': ('11020101', 'ROU-OFFICE RENTAL'),
            '2-7100': ('22010101', 'LEASE LIAB-NON CURRENT-OFFICE RENTAL'),
            # Specific User Requests
            'Domestic Boarding': ('75100000', 'TRAVEL-AIRTICKET'),
        }

    def get_sheet_data(self, excel_bytes):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
            sheet = wb.active
            return list(sheet.iter_rows(values_only=True))
        except Exception:
            return []

    def process(self, sg_bytes, in_bytes, id_bytes, stage=1):
        # Extract rows
        sg_rows = self.get_sheet_data(sg_bytes)
        in_rows = self.get_sheet_data(in_bytes)
        id_rows = self.get_sheet_data(id_bytes)

        # Stage 1: Just extract raw items
        sg_raw = []
        for row in sg_rows[1:]: # skip header
            if row and row[0]:
                sg_raw.append({'Code': str(row[0]), 'Description': str(row[1]) if len(row)>1 and row[1] else ''})
        
        in_raw = []
        in_unique = set()
        for row in in_rows[1:]:
            if row and row[0]:
                name = str(row[0]).strip()
                # Light filtering of obvious headers if needed
                if name and name not in in_unique and name.lower() not in ['particulars', 'in inr (₹)', 'capital account']:
                    in_unique.add(name)
                    in_raw.append({'Particulars': name})
                    
        id_raw = []
        id_unique = set()
        for row in id_rows[1:]:
            if row and row[0]:
                code = str(row[0]).strip()
                name = str(row[1]).strip() if len(row)>1 and row[1] else code
                if code and code not in id_unique:
                    id_unique.add(code)
                    id_raw.append({'COA Code': code, 'Account Name': name})

        if stage == 1:
            return {
                "sg_raw": sg_raw,
                "in_raw": in_raw,
                "id_raw": id_raw
            }

        # Stage 2: Mappings
        india_mapping = []
        indo_mapping = []
        
        def map_account(code, name):
            # 1. Hardcoded Domestic Boarding and Loadging Exception
            if name == 'Domestic Boarding and Loadging':
                return '75100000 / 75110000', 'TRAVEL-AIRTICKET / TRAVEL-HOTEL', '⚠ HUMAN APPROVAL REQUIRED'
            
            # 2. Standard Hardcoded mapping
            match = self.coa_mapping.get(code) or self.coa_mapping.get(name)
            if match:
                return match[0], match[1], '✓ MAPPED'
                
            # 3. AI Simulation Fallback (Search SG Code list)
            matches = [x for x in sg_raw if x['Description'].strip().upper() == name.strip().upper() or x['Code'] == code]
            if len(matches) == 1:
                return matches[0]['Code'], matches[0]['Description'], '✨ AI RESOLVED (Review)'
            elif len(matches) > 1:
                return 'MULTIPLE', 'MULTIPLE FOUND', '⚠ REVIEW REQUIRED'
            else:
                return 'UNMAPPED', 'NOT FOUND', '⚠ REVIEW REQUIRED'

        # Map India
        for item in in_raw:
            name = item['Particulars']
            sg_code, sg_desc, status = map_account(name, name)
            india_mapping.append({
                'Local Account': name,
                'SG Master Code': sg_code,
                'SG Account Description': sg_desc,
                'Status': status
            })
            
        # Map Indonesia
        for item in id_raw:
            code = item['COA Code']
            name = item['Account Name']
            sg_code, sg_desc, status = map_account(code, name)
            indo_mapping.append({
                'Local COA': code,
                'Local Name': name,
                'SG Master Code': sg_code,
                'SG Account Description': sg_desc,
                'Status': status
            })

        # Generate a simple Excel output so the download button still works
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Consolidation Mappings"
        ws.append(["Entity", "Local Account", "SG Master Code", "SG Account Description", "Status"])
        for m in india_mapping:
            ws.append(["India", m['Local Account'], m['SG Master Code'], m['SG Account Description'], m['Status']])
        for m in indo_mapping:
            ws.append(["Indonesia", f"{m['Local COA']} - {m['Local Name']}", m['SG Master Code'], m['SG Account Description'], m['Status']])
            
        out = io.BytesIO()
        wb.save(out)
        excel_base64 = base64.b64encode(out.getvalue()).decode('utf-8')
            
        return {
            "india_mapping": india_mapping,
            "indo_mapping": indo_mapping,
            "excel_base64": excel_base64
        }
